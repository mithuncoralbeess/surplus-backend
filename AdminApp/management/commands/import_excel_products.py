import os
import sys
import re
import datetime
from decimal import Decimal, InvalidOperation
import openpyxl
from django.core.management.base import BaseCommand
from django.utils import timezone
from AdminApp.models import Product, ProductImage, VendorDetails, MainCategory, SubCategory

# Fix Windows console charmap encoding issues for international product names
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def parse_clean_decimal(val, default=Decimal("0.00")):
    if val is None:
        return default
    if isinstance(val, (int, float, Decimal)):
        try:
            return Decimal(str(val)).quantize(Decimal("0.01"))
        except (InvalidOperation, ValueError):
            return default
    s = str(val).strip()
    # Remove currency codes, symbols, commas, spaces
    s = re.sub(r"[^\d.]", "", s)
    if not s:
        return default
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return default


def parse_clean_int(val, default=0):
    if val is None:
        return default
    if isinstance(val, int):
        return max(0, val)
    if isinstance(val, float):
        return max(0, int(val))
    s = str(val).strip()
    s = re.sub(r"[^\d]", "", s)
    return int(s) if s else default


def parse_year(val):
    if val is None:
        return None
    if isinstance(val, int):
        if 1900 <= val <= 2100:
            return val
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.year
    s = str(val).strip()
    m = re.search(r"\b(19\d\d|20\d\d)\b", s)
    if m:
        return int(m.group(1))
    return None


def parse_date(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    s = str(val).strip()
    if not s or s.lower() in ["none", "nil", "n/a"]:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def parse_datetime(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        if timezone.is_naive(val):
            return timezone.make_aware(val)
        return val
    if isinstance(val, datetime.date):
        dt = datetime.datetime.combine(val, datetime.time.min)
        return timezone.make_aware(dt)
    s = str(val).strip()
    if not s or s.lower() in ["none", "nil", "n/a"]:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            dt = datetime.datetime.strptime(s, fmt)
            return timezone.make_aware(dt)
        except ValueError:
            continue
    return None


def parse_bool(val):
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    s = str(val).strip().lower()
    return s in ["yes", "true", "1", "y"]


class Command(BaseCommand):
    help = "Import confirmed products from DB Final mithun.xlsx linking to existing vendor accounts by email"

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            type=str,
            default="DB Final mithun.xlsx",
            help="Path to the Excel file",
        )
        parser.add_argument(
            "--sheet",
            type=str,
            default="Confirmed - 221",
            help="Name of the worksheet tab to process",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Simulate import without creating products in the database",
        )
        parser.add_argument(
            "--limit",
            type=int,
            default=None,
            help="Limit number of rows to process",
        )

    def handle(self, *args, **options):
        file_path = options["file"]
        sheet_name = options["sheet"]
        dry_run = options["dry_run"]
        limit = options["limit"]

        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"File not found: {file_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Loading '{file_path}', sheet '{sheet_name}'..."))
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            self.stderr.write(self.style.ERROR(f"Sheet '{sheet_name}' not found. Sheets: {wb.sheetnames}"))
            return

        sheet = wb[sheet_name]
        row_iter = sheet.iter_rows(values_only=True)
        first_row = next(row_iter, None)
        if not first_row:
            self.stderr.write(self.style.ERROR("Sheet is empty."))
            return

        headers = [str(c).strip() if c is not None else "" for c in first_row[:32]]
        self.stdout.write(f"Header columns ({len(headers)}): {headers}")

        # Map header names to column index
        col_idx = {h: i for i, h in enumerate(headers) if h}

        # Cache vendors in memory for fast O(1) matching
        vendors_by_email = {
            v.email.lower().strip(): v for v in VendorDetails.objects.all() if v.email
        }
        self.stdout.write(f"Cached {len(vendors_by_email)} existing vendors in memory.")

        # Cache categories and subcategories in memory
        cats_by_name = {c.name.lower().strip(): c for c in MainCategory.objects.all()}
        subcats_cache = {}
        for sc in SubCategory.objects.select_related("main_category").all():
            key = (sc.main_category_id, sc.name.lower().strip())
            subcats_cache[key] = sc

        # Temporarily disconnect live sync signals & external embedding APIs during bulk creation
        from django.db.models.signals import post_save
        from AdminApp.signals import product_post_save_sync, product_image_post_save_sync
        import AdminApp.semantic_search as sem_search

        orig_sync_emb = getattr(sem_search, "sync_product_embedding", None)
        sem_search.sync_product_embedding = lambda *a, **kw: (None, False)
        post_save.disconnect(product_post_save_sync, sender=Product)
        post_save.disconnect(product_image_post_save_sync, sender=ProductImage)

        try:
            self._execute_import(sheet, col_idx, vendors_by_email, cats_by_name, subcats_cache, dry_run, limit)
        finally:
            if orig_sync_emb:
                sem_search.sync_product_embedding = orig_sync_emb
            post_save.connect(product_post_save_sync, sender=Product)
            post_save.connect(product_image_post_save_sync, sender=ProductImage)

    def _execute_import(self, sheet, col_idx, vendors_by_email, cats_by_name, subcats_cache, dry_run, limit):
        created_count = 0
        skipped_count = 0
        error_count = 0
        empty_streak = 0
        row_num = 1

        row_iter = sheet.iter_rows(values_only=True)
        # Skip header
        next(row_iter, None)

        for r_vals in row_iter:
            row_num += 1
            if not r_vals or not any(x is not None and str(x).strip() != "" for x in r_vals):
                empty_streak += 1
                if empty_streak >= 5:
                    break
                continue
            empty_streak = 0

            if limit and created_count >= limit:
                break

            def get_col(name, default=""):
                idx = col_idx.get(name)
                if idx is None or idx >= len(r_vals):
                    return default
                val = r_vals[idx]
                return val if val is not None else default

            product_name = str(get_col("Product Name", "")).strip()
            if not product_name:
                continue

            email_val = str(get_col("Email", "")).strip().lower()
            vendor = vendors_by_email.get(email_val)
            if not vendor:
                self.stderr.write(
                    self.style.WARNING(
                        f"Row {row_num}: No vendor found for email '{email_val}'! Skipping '{product_name}'."
                    )
                )
                error_count += 1
                continue

            # Category resolution
            raw_cat = str(get_col("Category", "")).strip()
            raw_subcat = str(get_col("Subcategory", "")).strip()

            cat_key = raw_cat.lower()
            if cat_key in ["toys, hobbies & sports", "toys, hobbies and sports", "toys and sports"]:
                cat_key = "toys & sports"

            main_cat = cats_by_name.get(cat_key)
            if not main_cat and raw_cat:
                if not dry_run:
                    main_cat = MainCategory.objects.create(name=raw_cat)
                    cats_by_name[raw_cat.lower()] = main_cat
                    self.stdout.write(f"Created missing MainCategory: '{raw_cat}'")

            # Subcategory resolution
            sub_cat = None
            if main_cat and raw_subcat:
                sub_key = (main_cat.id, raw_subcat.lower())
                sub_cat = subcats_cache.get(sub_key)
                if not sub_cat and not dry_run:
                    sub_cat = SubCategory.objects.create(main_category=main_cat, name=raw_subcat)
                    subcats_cache[sub_key] = sub_cat
                    self.stdout.write(f"Created missing SubCategory: '{raw_subcat}' under '{main_cat.name}'")

            brand_name = str(get_col("Brand", "")).strip()
            model_no = str(get_col("Model / Part No", "")).strip()
            raw_msrp = get_col("MSRP", None)
            raw_liq = get_col("Liquidation Price", 0.0)
            currency = str(get_col("Currency", "USD")).strip().upper()[:10] or "USD"
            quantity = parse_clean_int(get_col("Quantity", 0))
            description = str(get_col("Description", "")).strip()
            mfg_country = str(get_col("Manufacturing Country", "")).strip()
            mfg_year = parse_year(get_col("Manufacturing Year", None))
            warranty = str(get_col("Warranty", "")).strip()
            dimensions = str(get_col("Dimensions", "")).strip()
            expiry_date = parse_date(get_col("Expiry Date", None))
            third_party_cert = parse_bool(get_col("Third Party Certificate", False))
            documents = str(get_col("Documents", "")).strip()
            if documents.lower() in ["none", "nil", "n/a"]:
                documents = ""
            reason_to_sell = str(get_col("Reason to Sell", "")).strip()
            is_featured = parse_bool(get_col("Special / Featured", False))
            upload_batch = str(get_col("Upload Batch", "")).strip()
            if upload_batch.lower() in ["none", "nil", "n/a"]:
                upload_batch = "Excel Batch Confirmed-221"

            date_approved_dt = parse_datetime(get_col("Date Approved", None)) or timezone.now()
            excluded_countries_val = get_col("Excluded Countries", None)
            excluded_countries = []
            if isinstance(excluded_countries_val, list):
                excluded_countries = excluded_countries_val
            elif excluded_countries_val and str(excluded_countries_val).strip().lower() not in ["none", "nil", "n/a"]:
                excluded_countries = [c.strip() for c in str(excluded_countries_val).split(",") if c.strip()]

            liquidating_price = parse_clean_decimal(raw_liq, Decimal("0.00"))
            msrp = parse_clean_decimal(raw_msrp, None) if raw_msrp is not None else None

            # Calculate offer discount percentage if MSRP > liquidating_price
            offer = Decimal("0.00")
            if msrp and liquidating_price and msrp > liquidating_price:
                try:
                    pct = ((msrp - liquidating_price) / msrp) * Decimal("100.00")
                    offer = pct.quantize(Decimal("0.01"))
                except Exception:
                    offer = Decimal("0.00")

            # Inventory location defaults from vendor
            inventory_country = vendor.country or "United Arab Emirates"
            inventory_location = vendor.business_location or vendor.country or ""
            inventory_city = ""

            # Check if product already exists (idempotency check)
            existing_prod = Product.objects.filter(
                vendor=vendor,
                product_name=product_name,
                model_no=model_no,
            ).first()

            if existing_prod:
                skipped_count += 1
                continue

            if dry_run:
                self.stdout.write(
                    f"[DRY-RUN] Would create product: '{product_name}' (Model: '{model_no}') | "
                    f"Vendor: {vendor.vendor_id} ({vendor.email}) | Cat: {main_cat} | Price: {currency} {liquidating_price}"
                )
                created_count += 1
                continue

            # Create Product
            product = Product.objects.create(
                vendor=vendor,
                admin=None,
                product_name=product_name,
                category=main_cat,
                subcategory=sub_cat,
                brand_name=brand_name,
                model_no=model_no,
                manufacturing_country=mfg_country,
                inventory_country=inventory_country,
                inventory_location=inventory_location,
                inventory_city=inventory_city,
                intake_channel=upload_batch,
                manufacturing_year=mfg_year,
                dimensions=dimensions,
                expiry_date=expiry_date,
                excluded_countries=excluded_countries,
                quantity=quantity,
                currency=currency,
                liquidating_price=liquidating_price,
                msrp=msrp,
                offer=offer,
                description=description,
                reason_to_sell=reason_to_sell,
                warranty=warranty,
                third_party_certificate=third_party_cert,
                third_party_documents=documents,
                enquiry_status="APPROVED",
                is_active=True,
                date_approved=date_approved_dt,
                is_featured=is_featured,
            )

            # Images creation
            feat_img_url = str(get_col("Featured Image URL", "")).strip()
            all_imgs_raw = str(get_col("All Image URLs", "")).strip()

            created_image_urls = set()
            if feat_img_url and feat_img_url.lower() not in ["none", "nil", "n/a"]:
                ProductImage.objects.create(
                    product=product,
                    image_url=feat_img_url,
                    is_real_photo=True,
                )
                created_image_urls.add(feat_img_url)

            if all_imgs_raw and all_imgs_raw.lower() not in ["none", "nil", "n/a"]:
                urls = [u.strip() for u in all_imgs_raw.split(",") if u.strip()]
                for u in urls:
                    if u not in created_image_urls:
                        ProductImage.objects.create(
                            product=product,
                            image_url=u,
                            is_real_photo=True,
                        )
                        created_image_urls.add(u)

            created_count += 1
            if created_count % 25 == 0 or created_count == 221:
                self.stdout.write(
                    self.style.SUCCESS(
                        f"[{created_count}/221] Created {product.product_id}: '{product.product_name}' "
                        f"(Vendor: {vendor.vendor_id}) | Images: {len(created_image_urls)}"
                    )
                )

        mode_str = "[DRY-RUN] " if dry_run else ""
        self.stdout.write(
            self.style.SUCCESS(
                f"\n{mode_str}Complete! Processed {row_num - 1} rows. "
                f"Created: {created_count}, Skipped (Already existed): {skipped_count}, Errors: {error_count}."
            )
        )
